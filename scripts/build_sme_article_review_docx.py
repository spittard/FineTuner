"""
Build a single Word document for SME review: instructions + yea/nea + full article draft.
Requires: pip install python-docx

Usage (from repo root):
  python scripts/build_sme_article_review_docx.py
Writes: docs/ARTICLE_SME_REVIEW_EDITABLE.docx
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt
except ImportError:
    print("Install: pip install python-docx", file=sys.stderr)
    sys.exit(1)

REPO = Path(__file__).resolve().parents[1]
ARTICLE_MD = REPO / "docs" / "ARTICLE_HIGH_STAKES_COMPANY_MATCHING_SME.md"
OUT_DOCX = REPO / "docs" / "ARTICLE_SME_REVIEW_EDITABLE.docx"


def _strip_inline_md(s: str) -> str:
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"\*(.+?)\*", r"\1", s)
    s = s.replace("`", "")
    return s


def _sanitize_publication_notes(text: str) -> str:
    return text.replace(
        "Company Matching at Scale Without Gaslighting Your SMEs",
        "Company Matching at Scale While Protecting SME Time",
    )


def add_markdown_body(doc: Document, md_text: str) -> None:
    md_text = _sanitize_publication_notes(md_text)
    for raw_line in md_text.splitlines():
        line = raw_line.rstrip()
        if not line or line.strip() == "---":
            continue
        if line.startswith("### "):
            doc.add_heading(_strip_inline_md(line[4:]), level=3)
        elif line.startswith("## "):
            doc.add_heading(_strip_inline_md(line[3:]), level=2)
        elif line.startswith("# "):
            doc.add_heading(_strip_inline_md(line[2:]), level=1)
        elif line.startswith("- "):
            p = doc.add_paragraph(_strip_inline_md(line[2:]), style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
        else:
            p = doc.add_paragraph(_strip_inline_md(line))
            p.paragraph_format.space_after = Pt(4)


def main() -> None:
    if not ARTICLE_MD.is_file():
        print(f"Missing {ARTICLE_MD}", file=sys.stderr)
        sys.exit(1)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    doc.add_heading("SME review — public article (editable)", level=0)

    doc.add_heading("Purpose", level=1)
    doc.add_paragraph(
        "This Word file is for you to read, edit if you want, and record a clear go / no-go. "
        "Nothing here names clients, brands, or individuals. The article draft is generic "
        "(hospitality-adjacent / events / procurement context only)."
    )

    doc.add_heading("Your decision", level=1)
    doc.add_paragraph("Check one and add your name and date below.")
    doc.add_paragraph("Yea — I am comfortable with publication of the article draft in the Appendix (possibly after my edits in this file).", style="List Bullet")
    doc.add_paragraph("Nea — I am not comfortable; see notes below.", style="List Bullet")
    doc.add_paragraph("")
    doc.add_paragraph("Name: ___________________________    Date: ___________________________")

    doc.add_heading("Optional notes (edit freely)", level=1)
    doc.add_paragraph(
        "Lines or themes to change, titles that feel off, anything that could be read as "
        "blaming people instead of broken process — write here:"
    )
    for _ in range(5):
        doc.add_paragraph("________________________________________________________________________________")

    doc.add_heading("Quick reminder — what the article argues", level=1)
    for t in (
        "SMEs do high-stakes “plugging” work; bad tooling shifts burden to rework and erodes trust.",
        "Semantic search alone is not enough; hybrid ranking + geo + scores need discipline.",
        "Machine assessment covers every row for triage; it does not replace human judgment.",
        "Worst cases become regression tests before expensive full reruns.",
        "Honest scope beats claiming “everyone read everything” without evidence.",
    ):
        doc.add_paragraph(t, style="List Bullet")

    doc.add_page_break()
    doc.add_heading("Appendix — article draft (for Medium-style publication)", level=1)
    doc.add_paragraph(
        "You may edit directly in Word. For version control, consider copying changes back to "
        f"`docs/ARTICLE_HIGH_STAKES_COMPANY_MATCHING_SME.md` if the team uses that file as source."
    )
    doc.add_paragraph("")

    article = ARTICLE_MD.read_text(encoding="utf-8")
    add_markdown_body(doc, article)

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    print(f"Wrote {OUT_DOCX}")


if __name__ == "__main__":
    main()
